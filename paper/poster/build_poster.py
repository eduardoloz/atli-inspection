#!/usr/bin/env python3
"""Build a 48x36in research poster (PPTX) from the paper's verified content.

Every number matches paper/main.tex, which traces to results/*.md.
Import into Google Slides: Drive > New > File upload > open with Google Slides.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pathlib import Path

HERE = Path(__file__).resolve().parent
FIGS = HERE.parent / "figures"

NAVY = RGBColor(0x1F, 0x38, 0x5D)
BLUE = RGBColor(0x21, 0x66, 0xAC)
GREEN = RGBColor(0x1A, 0x98, 0x50)
RED = RGBColor(0xB0, 0x3A, 0x2E)
BOX_BG = RGBColor(0xF4, 0xF6, 0xF8)
DARK = RGBColor(0x21, 0x21, 0x21)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(48)
prs.slide_height = Inches(36)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


def rect(x, y, w, h, fill, rounded=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def textbox(x, y, w, h, lines, align=PP_ALIGN.LEFT):
    """lines: list of (text, size_pt, bold, color, space_after_pt)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color, space) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        run = p.add_run()
        run.text = text
        f = run.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
    return tb


def section(x, y, w, h, title, color=BLUE):
    rect(x, y, w, h, BOX_BG)
    textbox(x + 0.45, y + 0.3, w - 0.9, 1.0, [(title, 30, True, color, 0)])
    return y + 1.45  # content start y


def body(x, y, w, items, size=19):
    lines = [("• " + t, size, False, DARK, 8) for t in items]
    textbox(x + 0.45, y, w - 0.9, 10, lines)


def picture(x, y, w, path):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))


def caption(x, y, w, text):
    textbox(x, y, w, 0.6, [(text, 15, False, RGBColor(0x66, 0x66, 0x66), 0)],
            align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- header band
rect(0, 0, 48, 4.6, NAVY, rounded=False)
textbox(1.2, 0.55, 45.6, 2.2, [(
    "Data-Centric Rare-Defect Detection for Autonomous Transmission-Line "
    "Inspection with Nano-Scale Detectors", 62, True, WHITE, 0)], align=PP_ALIGN.CENTER)
textbox(1.2, 3.15, 45.6, 0.9, [(
    "Summer 2026 research campaign  ·  Dept. of Electrical & Computer Engineering, "
    "University of Nevada, Las Vegas  ·  paper draft in preparation (author list pending)",
    24, False, RGBColor(0xD5, 0xDE, 0xEC), 0)], align=PP_ALIGN.CENTER)

X = [0.9, 16.55, 32.2]
W = 14.9

# ---------------------------------------------------------------- column 1
cy = section(X[0], 5.0, W, 11.4, "1 · The problem: one rare, safety-critical class")
body(X[0], cy, W, [
    "UAV + deep-learning line inspection is data-starved: of 73 surveyed studies, only 23% use public datasets and 78% flag small-object difficulty.",
    "ATLI benchmark: 732 UAV images, 7 classes, defect-spot labels. Defective Damper: 113 instances vs. 1,460 Normal Dampers (12.9:1).",
    "Damper defects are the worst or near-worst class in every multi-class study we reviewed (DSA-Net 0.636, TLDD 0.40, HC-ViT AP75 0.371).",
    "Question: what actually raises rare-class AP — more data, more resolution, more augmentation, or more model?",
])

cy = section(X[0], 16.9, W, 15.1, "2 · Evaluation hazards nobody audits", RED)
body(X[0], cy, W, [
    "Provenance audit (pHash + filenames): public datasets duplicate INTO the evaluation split — CPLID: 249 images inside ATLI (37 in test); DVDI: ~50% overlap; PTL-AI Furnas: 32 exact val/test duplicates.",
    "“Adding more public data” can silently train on your test set. Every external experiment here passes a mandatory leakage gate.",
    "Single-split rare-class AP swings 0.66–0.88 across equally valid draws — single-run numbers are close to uninformative.",
    "Our protocol: proper 5-fold cross-validation (val ≠ test) + seed-averaged means ± std; decontaminated 797-image clean benchmark rebuilt (recipe margin survives: 0.784 vs 0.736 mAP).",
])

# ---------------------------------------------------------------- column 2
cy = section(X[1], 5.0, W, 11.4, "3 · What failed: external data, at every dose", RED)
body(X[1], cy, W, [
    "Community damper data mixed at ratios 1:1 → 11.7:1: no gain (DD AP 0.58–0.65 vs baseline 0.641); recall drops, precision holds.",
    "In-distribution control: ≈0.91 AP on the same external data → the labels are learnable; the transfer fails. What matters is label-style consistency, not domain proximity.",
    "In-domain source pretraining: −7.5 mAP (genuine negative transfer). Scale-matched curation, staged pretraining: no gain.",
    "Rotation augmentation — standard in the field — hurts axis-aligned detection (10°: −0.10 DD AP; 45°: −0.037 mAP).",
])

cy = section(X[1], 16.9, W, 15.1, "4 · What worked: a native-only recipe", GREEN)
body(X[1], cy, W, [
    "1280 px training + 3× oversampling of damper-bearing images + scale-0.9 augmentation. Zero new labels, zero external data.",
    "All three nano models gain +4–6 mAP and +6–9 DD AP points under 5-fold CV. YOLOv11n (5-fold means ± std):",
], size=19)
# mini results table
tbl_y = cy + 4.4
tblshape = slide.shapes.add_table(3, 4, Inches(X[1] + 0.45), Inches(tbl_y),
                                  Inches(W - 0.9), Inches(2.6))
tbl = tblshape.table
hdr = ["YOLOv11n", "mAP@0.5", "DD AP", "DD recall"]
rows = [["base (640)", "0.722 ± 0.019", "0.743 ± 0.074", "0.695"],
        ["champ (recipe)", "0.785 ± 0.023", "0.812 ± 0.080", "0.764"]]
for j, t in enumerate(hdr):
    c = tbl.cell(0, j); c.text = t
for i, r in enumerate(rows):
    for j, t in enumerate(r):
        tbl.cell(i + 1, j).text = t
for r in range(3):
    for j in range(4):
        cell = tbl.cell(r, j)
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(18)
                run.font.name = FONT
                run.font.bold = (r == 0) or (r == 2 and j == 0)
from PIL import Image as _Img
_pc = _Img.open(FIGS / "fig_perclass_clean.png")
_pc_h = (W - 1.0) / (_pc.width / _pc.height)
picture(X[1] + 0.5, tbl_y + 3.0, W - 1.0, FIGS / "fig_perclass_clean.png")
caption(X[1], tbl_y + 3.0 + _pc_h + 0.2, W,
        "Per-class AP on the decontaminated benchmark (3 seeds; error bars = std).")

# ---------------------------------------------------------------- column 3
cy = section(X[2], 5.0, W, 13.5, "5 · A 2.6M nano model matches a 47M transformer")
fw = 12.4
fx = X[2] + (W - fw) / 2
picture(fx, cy + 0.1, fw, FIGS / "fig_frontier.png")
fh = fw / (1050 / 750)
body(X[2], cy + fh + 0.4, W, [
    "Identical folds, MMDetection baselines (stock configs): YOLOv11n + recipe ties DINO-4scale on mAP@0.5 (0.785 both) and is comparable on the rare class (0.812 vs 0.789, within fold noise) at ~18× fewer parameters.",
], size=19)

cy = section(X[2], 18.9, W, 13.1, "6 · It exports to the edge", GREEN)
fw2 = 10.6
picture(X[2] + (W - fw2) / 2, cy + 0.05, fw2, FIGS / "fig_deploy.png")
fh2 = fw2 / (1050 / 780)
body(X[2], cy + fh2 + 0.3, W, [
    "768 px retrain → TensorRT FP16: 8.1 MB engine, −0.009 mAP export cost, fits a Jetson Nano's 2 GB envelope. Throughput (~18 fps) is projected, not yet measured on-device.",
], size=19)

# ---------------------------------------------------------------- takeaway band
rect(0.9, 32.4, 46.2, 2.9, NAVY)
textbox(1.5, 32.65, 10.5, 2.4, [("Takeaways", 30, True, WHITE, 0)])
textbox(12.4, 32.62, 34.3, 2.6, [
    ("1  Exhaust native data-centric levers (resolution, oversampling, scale aug) before harvesting external data.        "
     "2  pHash-gate every external source — public inspection datasets recycle each other.", 19.5, False, WHITE, 6),
    ("3  Report rare-class results as cross-validated / multi-seed means — single splits swing by 0.2+.        "
     "4  Next: on-device latency measurement, class-weighted losses, larger damper test pool.", 19.5, False, WHITE, 0),
])

out = HERE / "ATLI_poster_48x36.pptx"
prs.save(out)
print("saved", out)
# structural sanity
p2 = Presentation(str(out))
print("slides:", len(p2.slides), "| shapes:", len(p2.slides[0].shapes._spTree))
